"""Read Dr. Aishu's consolidated Excel workbook.

The workbook is built by `scripts/build_dataset_workbook.py` from the raw
downloads, so it is a derived artefact rather than hand-edited data. This module
opens it lazily and caches the parsed sheets, and it degrades gracefully: if the
workbook has not been built, `available()` is False and every lookup returns an
empty result instead of raising, so the server still starts on a fresh clone.

openpyxl is a build/analysis dependency, not a server one. It is imported inside
the loader so the app keeps running when only the runtime requirements are
installed; the endpoints then report the dataset as unavailable.
"""
from __future__ import annotations

import threading
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_PATH = ROOT / "data" / "external" / "dr_aishu_dataset.xlsx"

# Sheets the workbook is expected to carry, and the columns each must have. The
# builder writes these; the API and tests read them by name.
REQUIRED_SHEETS: dict[str, tuple[str, ...]] = {
    "Diseases": ("disease_id", "name", "in_kb", "treatments", "sources"),
    "Symptoms": ("symptom_id", "label", "in_kb"),
    "Disease_Symptoms": ("disease", "symptom"),
    "Precautions": ("disease", "precaution_1"),
    "Images": ("image_id", "diagnosis", "modality", "source", "license"),
    "MedicineNet": ("condition", "slug", "url"),
    "Sources": ("id", "url", "license", "status"),
}

# Sheets the builder writes that are useful but not required: an older workbook
# without them still loads, so `available()` stays a check on the core contract.
OPTIONAL_SHEETS: dict[str, tuple[str, ...]] = {
    "DeepLesion": ("file_name", "coarse_lesion_type", "patient_age", "split", "license"),
    "Dataset_Catalogue": ("study_refs", "disease", "dataset", "url", "license"),
}

_cache: dict[str, Any] = {"loaded": False, "path": None, "sheets": {}, "error": None}
_lock = threading.Lock()


def _read(path: Path) -> dict[str, list[dict]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        out: dict[str, list[dict]] = {}
        for name in wb.sheetnames:
            ws = wb[name]
            rows = ws.iter_rows(values_only=True)
            try:
                header = [str(c).strip() if c is not None else "" for c in next(rows)]
            except StopIteration:
                out[name] = []
                continue
            records = []
            for raw in rows:
                if raw is None or all(c is None for c in raw):
                    continue
                records.append({
                    header[i]: (raw[i] if i < len(raw) else None)
                    for i in range(len(header))
                    if header[i]
                })
            out[name] = records
        return out
    finally:
        wb.close()


def load(path: Path | None = None, force: bool = False) -> dict[str, list[dict]]:
    """Parse the workbook once and cache it. Returns sheet name -> rows."""
    target = Path(path) if path else DEFAULT_PATH
    with _lock:
        if _cache["loaded"] and _cache["path"] == target and not force:
            return _cache["sheets"]
        sheets: dict[str, list[dict]] = {}
        error = None
        if target.exists():
            try:
                sheets = _read(target)
            except Exception as exc:  # noqa: BLE001 - reported, never fatal
                error = f"{type(exc).__name__}: {exc}"
        _cache.update({"loaded": True, "path": target, "sheets": sheets, "error": error})
        return sheets


def available(path: Path | None = None) -> bool:
    """True when a workbook exists and parsed into the expected sheets."""
    sheets = load(path)
    return bool(sheets) and all(name in sheets for name in REQUIRED_SHEETS)


def reset() -> None:
    """Drop the cache. Tests use this after writing a temporary workbook."""
    with _lock:
        _cache.update({"loaded": False, "path": None, "sheets": {}, "error": None})


def sheet(name: str, path: Path | None = None) -> list[dict]:
    return load(path).get(name, [])


def status(path: Path | None = None) -> dict[str, Any]:
    """Counts and health, for the API and the dataset page."""
    target = Path(path) if path else DEFAULT_PATH
    sheets = load(target)
    ok = available(target)
    counts = {name: len(sheets.get(name, [])) for name in REQUIRED_SHEETS}
    for name in OPTIONAL_SHEETS:
        if name in sheets:
            counts[name] = len(sheets.get(name, []))
    missing = [n for n in REQUIRED_SHEETS if n not in sheets]
    missing_columns = {
        n: [c for c in cols if not sheets.get(n) or c not in (sheets.get(n) or [{}])[0]]
        for n, cols in REQUIRED_SHEETS.items()
        if sheets.get(n)
    }
    return {
        "available": ok,
        "path": str(target.relative_to(ROOT)) if target.exists() else None,
        "bytes": target.stat().st_size if target.exists() else 0,
        "sheets": counts,
        "missing_sheets": missing,
        "missing_columns": {k: v for k, v in missing_columns.items() if v},
        "error": _cache.get("error"),
        "hint": None if ok else "Run: python scripts/fetch_sources.py && python scripts/build_dataset_workbook.py",
    }


def diseases(path: Path | None = None, in_kb_only: bool = False) -> list[dict]:
    rows = sheet("Diseases", path)
    if in_kb_only:
        rows = [r for r in rows if (r.get("in_kb") or "").lower() == "yes"]
    return rows


def find_disease(name_or_id: str, path: Path | None = None) -> dict | None:
    """Look a condition up by id or name, case- and punctuation-insensitively."""
    import re

    def key(text: str) -> str:
        return re.sub(r"[^a-z0-9]", "", (text or "").lower())

    want = key(name_or_id)
    if not want:
        return None
    for row in diseases(path):
        if key(row.get("disease_id", "")) == want or key(row.get("name", "")) == want:
            return row
    return None


def images(path: Path | None = None, modality: str | None = None) -> list[dict]:
    rows = sheet("Images", path)
    if modality:
        rows = [r for r in rows if (r.get("modality") or "").lower() == modality.lower()]
    return rows


def image_label_counts(path: Path | None = None) -> dict[str, dict[str, int]]:
    """How many images per modality and class, for the dataset page."""
    out: dict[str, dict[str, int]] = {}
    for row in sheet("Images", path):
        mod = (row.get("modality") or "unknown").lower()
        cls = (row.get("diagnosis") or "unknown")
        out.setdefault(mod, {})
        out[mod][cls] = out[mod].get(cls, 0) + 1
    return out


def lesion_type_counts(path: Path | None = None) -> dict[str, int]:
    """How many DeepLesion annotations fall in each coarse lesion type.

    The type is a numeric code in the source CSV; the count is what the dataset
    page shows, so the codes are returned as-is rather than guessed at.
    """
    out: dict[str, int] = {}
    for row in sheet("DeepLesion", path):
        code = str(row.get("coarse_lesion_type") or "unknown")
        out[code] = out.get(code, 0) + 1
    return dict(sorted(out.items(), key=lambda kv: -kv[1]))


def sources(path: Path | None = None) -> list[dict]:
    return sheet("Sources", path)


def medicinenet_conditions(path: Path | None = None, limit: int | None = None) -> list[dict]:
    rows = sheet("MedicineNet", path)
    return rows[:limit] if limit else rows


def summary(path: Path | None = None) -> dict[str, Any]:
    """A compact view for the home/insights page."""
    st = status(path)
    if not st["available"]:
        return st
    return {
        **st,
        "image_labels": image_label_counts(path),
        "conditions_in_kb": sum(
            1 for r in diseases(path) if (r.get("in_kb") or "").lower() == "yes"
        ),
        "licenses": sorted({(s.get("license") or "unknown") for s in sources(path)}),
    }
